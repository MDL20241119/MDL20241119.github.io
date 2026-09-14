#!/usr/bin/env python3
"""Normalize Fukuoka public facility, railway and observed data. Source rows retained."""
import sys,json,csv,io,zipfile,hashlib,re,collections
from pathlib import Path
import openpyxl
from shapely.geometry import shape,Point,mapping
from shapely.ops import unary_union
RAW=Path(sys.argv[1]);ROOT=Path(__file__).resolve().parents[1]/'fukuoka-mobility'
TODAY='2026-09-14'
def write(name,obj):
 p=ROOT/name;p.parent.mkdir(parents=True,exist_ok=True);p.write_text(json.dumps(obj,ensure_ascii=False,separators=(',',':'))+'\n')
def code(c):
 c=str(c)[:5]
 return '40100' if c.startswith('4010') else '40130' if c.startswith('4013') else c
bounds=json.loads((ROOT/'gaps/data/municipalities.geojson').read_text())['features'];geoms=[(f['properties'],shape(f['geometry'])) for f in bounds];names={p['code']:p['name'] for p,g in geoms};whole=unary_union([g for p,g in geoms])
def municipality(lat,lon):
 p=Point(lon,lat)
 return next((a for a,g in geoms if g.covers(p)),None)
def num(v):
 try:return float(str(v).replace(',',''))
 except:return None
def csvrows(b):
 for enc in ['utf-8-sig','cp932','utf-16']:
  try:return list(csv.DictReader(io.StringIO(b.decode(enc))))
  except UnicodeError:pass
 raise ValueError('CSV encoding')
def pick(r,*keys):return next((str(r[k]).strip() for k in keys if r.get(k) is not None and str(r[k]).strip()),'')
places=[];rejected=[]
z=zipfile.ZipFile(RAW/'hospital.zip')
for rownum,r in enumerate(csvrows(z.read(z.namelist()[0])),2):
 if r['都道府県コード']!='40':continue
 lat,lon=num(r['所在地座標（緯度）']),num(r['所在地座標（経度）']);c=code('40'+r['市区町村コード'].zfill(3))
 if lat is None or lon is None or c not in names or not (32.7<lat<34.4 and 129.9<lon<131.5):rejected.append({'source':'厚生労働省','name':r['正式名称'],'reason':'有効な座標または市町村コードなし'});continue
 places.append(dict(id='hospital-'+r['ID'],name=r['正式名称'],category='hospital',city=names[c],municipalityCode=c,lat=lat,lon=lon,address=r['所在地'],url=r['案内用ホームページアドレス'] or None,source='厚生労働省 医療情報ネット',sourceUrl='https://www.mhlw.go.jp/stf/seisakunitsuite/bunya/kenkou_iryou/iryou/newpage_43373.html',sourceResourceUrl='https://www.mhlw.go.jp/content/11121000/01-1_hospital_facility_info_20260601.csv.zip',sourceRow=rownum,dataAsOf='2026-06-01',retrievedAt=TODAY,license='公共データ利用規約 第1.0版（PDL1.0）',licenseUrl='https://www.digital.go.jp/resources/open_data/public_data_license_v1.0',openingHours=None,wheelchairAccess=None,dataNote='医療機関が報告した所在地。診療の受付時間・入口・当日の利用条件は未確認。'))
seen=set();manifest=json.loads((RAW/'public-manifest.json').read_text())
for m in manifest:
 p=RAW/m['file']
 if not p.exists():continue
 try:rows=csvrows(p.read_bytes())
 except Exception as e:rejected.append({'source':m['sourceUrl'],'reason':str(e)});continue
 for rownum,r in enumerate(rows,2):
  name=pick(r,'名称','施設名');lat=num(pick(r,'緯度'));lon=num(pick(r,'経度'))
  if not name:continue
  if lat is None or lon is None:rejected.append({'source':m['sourceUrl'],'name':name,'reason':'座標なし'});continue
  area=municipality(lat,lon)
  if not area:rejected.append({'source':m['sourceUrl'],'name':name,'reason':'福岡県表示行政界外・位置要確認'});continue
  dedup=(name,round(lat,4),round(lon,4))
  if dedup in seen:continue
  seen.add(dedup)
  category='school' if re.search('学校|大学|高等|学園',name) else 'library' if '図書' in name else 'clinic' if re.search('診療所|クリニック',name) else 'hospital' if '病院' in name else 'civic'
  address=pick(r,'所在地_連結表記','住所','所在地') or ''.join(pick(r,k) for k in ['所在地_都道府県','所在地_市区町村','所在地_町字','所在地_番地以下'])
  if not address:address='福岡県'+area['name']
  places.append(dict(id='public-'+hashlib.sha256((m['file']+str(rownum)).encode()).hexdigest()[:12],name=name,category=category,city=area['name'],municipalityCode=area['code'],lat=lat,lon=lon,address=address,url=pick(r,'URL') or None,source=m['provider']+' '+m['title'],sourceUrl=m['sourceUrl'],sourceResourceUrl=m['url'],sourceRow=rownum,dataAsOf=None,sourceUpdatedAt=m.get('sourceUpdatedAt'),retrievedAt=TODAY,license=m['license'],licenseUrl=m['licenseUrl'],openingHours=None,wheelchairAccess=None,dataNote='自治体公開CSVを名称・座標で重複除去し分類。公開更新日は施設の現況確認日ではありません。'))
# Public facts from the operator's own store locator; no logos, photos or descriptions.
shopping=[]
for r in json.loads((RAW/'aeon-shops.json').read_text())['shops']:
 if not r['address'].startswith('福岡県') or not re.match('イオン(?!バイク)|マックスバリュ|ザ・ビッグ|レッドキャベツ',r['nameKanji']):continue
 lat,lon=num(r['latitude']),num(r['longitude'])
 if lat is None or lon is None:continue
 area=municipality(lat,lon)
 if not area:continue
 u='https://tenpo.aeon-kyushu.info/detail/'+r['storeCode']+'/'
 times=r.get('businessHours') or [];dayNames={'MONDAY':'月','TUESDAY':'火','WEDNESDAY':'水','THURSDAY':'木','FRIDAY':'金','SATURDAY':'土','SUNDAY':'日'}
 hours=' / '.join(dayNames.get(t['name'],t['name'])+' '+str(t.get('openTime',''))[:5]+'〜'+str(t.get('closeTime',''))[:5] for t in times)
 shopping.append(dict(id='store-'+str(r['id']),name=r['nameKanji'],category='shopping',city=area['name'],municipalityCode=area['code'],lat=lat,lon=lon,address=r['address'],url=u,source='イオン九州 店舗公式案内',sourceUrl=u,sourceResourceUrl='https://tenpo.aeon-kyushu.info/',dataAsOf=None,retrievedAt=TODAY,checkedAt=TODAY,license=None,licenseUrl=None,openingHours=hours or None,wheelchairAccess=None,needsCoordinateConfirmation=True,dataNote='運営者の公開店舗検索から所在地・営業時間を整理。掲載施設の代表位置。入口・当日の営業は個別確認。県内店舗の網羅ではありません。'))
write('access/destinations.json',places);write('access/shopping.json',shopping)
write('data/facility-audit.json',{'retrievedAt':TODAY,'publicFacilities':len(places),'shopping':len(shopping),'excluded':rejected,'sourceFiles':[dict(file=m['file'],sourceUrl=m['sourceUrl'],url=m['url'],sha256=hashlib.sha256((RAW/m['file']).read_bytes()).hexdigest()) for m in manifest if (RAW/m['file']).exists()]})
# Railway station and line geometries are from the national dataset, clipped to Fukuoka.
z=zipfile.ZipFile(RAW/'rail.zip');rail={}
for kind,suffix in [('stations','Station.geojson'),('lines','RailroadSection.geojson')]:
 src=json.loads(z.read(next(n for n in z.namelist() if n.endswith(suffix))))
 fs=[]
 for f in src['features']:
  g=shape(f['geometry'])
  if not g.intersects(whole):continue
  clipped=g.intersection(whole)
  if clipped.is_empty:continue
  fs.append({'type':'Feature','properties':f['properties'],'geometry':mapping(clipped)})
 rail[kind]={'type':'FeatureCollection','features':fs}
rail['meta']={'sourceUrl':'https://nlftp.mlit.go.jp/ksj/gml/datalist/KsjTmplt-N02-2025.html','dataAsOf':'2025-12-31','retrievedAt':TODAY,'license':'CC BY 4.0','licenseUrl':'https://nlftp.mlit.go.jp/ksj/other/agreement_01.html','attribution':'国土数値情報（鉄道データ・2025年）（国土交通省）を加工','note':'鉄道施設の位置・形状。時刻表・停車便・運行状況は未収録。駅施設の線分数でありユニーク駅数ではない。'}
write('data/rail.json',rail)
z=zipfile.ZipFile(RAW/'busstops.zip');bus=json.loads(z.read(next(n for n in z.namelist() if n.endswith('.geojson'))))
for feature in bus['features']:
    feature['properties']={key:value for key,value in feature['properties'].items() if value is not None}
    feature['geometry']['coordinates']=[round(value,6) for value in feature['geometry']['coordinates']]
bus['metadata']={'processing':'原典のnull属性キーを省略。省略属性は未収録を意味します。座標は小数点以下6桁に丸めています。','source':'https://nlftp.mlit.go.jp/ksj/gml/data/P11/P11-22/P11-22_40_SHP.zip'}
write('data/bus-stop-inventory.geojson',bus)
# Official annual average daily subway boardings. Preserve published row totals, do not recompute rounding.
url='https://data.bodik.jp/dataset/401307_tikatetuekibetujoushajinin'
wb=openpyxl.load_workbook(RAW/'facilities/401307_tikatetuekibetujoushajinin.xlsx',data_only=True);ws=wb.active;years=[(c,ws.cell(2,c).value) for c in range(2,ws.max_column+1) if isinstance(ws.cell(2,c).value,int)]
stops=[];history=[];routes=[]
for row in range(3,48):
 name=str(ws.cell(row,1).value or '').strip();line='空港線' if row<=17 else '箱崎線' if row<=26 else '空港線・箱崎線' if row==27 else '七隈線' if row<=46 else '全線'
 for col,year in years:
  val=ws.cell(row,col).value
  if not isinstance(val,(float,int)):continue
  r=dict(year=year,stop=name,route=line,value=val,unit='人／日（年度平均乗車人員）',municipality='福岡市',source_url=url,source_cell=ws.cell(row,col).coordinate,source_sheet=ws.title)
  if row in [17,26,46,47]:
   history.append(r)
   if year==2024:routes.append(dict(route=line,operator='福岡市交通局',period='2024年度',passengers=val,revenue=None,scope='年度平均乗車人員・人／日',source_url=url))
  elif row!=27 and year==2024:stops.append(r)
financeUrl='https://www.city.fukuoka.lg.jp/gikaizimukyoku/giji/shisei/documents/20250926-koutsu-siryo.pdf'
finance=[dict(area='福岡市',service='福岡市地下鉄 全線',period='2024年度（2024-04-01〜2025-03-31）',revenue=36065495968,cost=29972712775,cost_label='営業費用（減価償却費を含む）',cost_per_passenger=None,passengers=191340519,transport_revenue=33859964608,operating_profit=6092783193,ordinary_profit=9030064111,net_profit=9316495094,subsidy=1104864209,depreciation=13315834735,source_url=financeUrl,source_page=39,unit='円・損益計算書（税抜）',note='営業収益と営業費用の比較。営業外・特別損益を含む純利益と混同しない。乗車人員は延べ人数。')]
expenses=[('線路保存費',3216207546),('電路保存費',2018654994),('車両保存費',2108477685),('運転費',2694248162),('運輸管理費',2052882201),('運輸費',3815420861),('研修所費',61473771),('一般管理費',689512820),('減価償却費',13315834735)]
assert sum(v for k,v in expenses)==finance[0]['cost']
ptUrl='https://www.pref.fukuoka.lg.jp/uploaded/life/796583_62719454_misc.pdf'
pt=[dict(mode=mode,share_pct=share,year=year,source_url=ptUrl,source_page=8,unit='構成比・公表丸め値％') for year,values in [(1993,[24,13,49,6,8]),(2005,[18,11,58,5,8]),(2017,[18,11,55,5,11])] for mode,share in zip(['徒歩','二輪車','自動車','バス','鉄道'],values)]
# Municipality census values match by full name; wards are not summed again.
ws2=openpyxl.load_workbook(RAW/'facilities/400009_kokuseichousa_r02.xlsx',data_only=True).active;pop=[]
for i,row in enumerate(ws2.values,1):
 text=''.join(str(v or '') for v in row[:6]).replace(' ','').replace('　','')
 for c,name in names.items():
  if text==name and isinstance(row[6],(float,int)):
   pop.append(dict(municipality=name,code=c,population=row[6],male=row[7],female=row[8],households=row[9],year=2020,source_url='https://data.bodik.jp/dataset/400009_kokuseichousa_r02',source_row=i))
analysis=dict(meta={'title':'福岡県 交通データ','retrievedAt':TODAY,'observedScope':'各原典の時点・対象範囲。福岡県の全交通需要・財務を網羅するものではない。'},route_actuals=routes,stop_actuals=stops,subway_history=history,finance=finance,expenses=[{'item':k,'yen':v} for k,v in expenses],pt_modes=pt,pt_scope='第5回北部九州圏PT・2017年。福岡県のほぼ全域と佐賀県鳥栖市・基山町。全交通手段・代表交通手段。福岡県単独の集計ではない。',municipal_population=pop,commute_od=[],sources=[{'title':'福岡市 地下鉄駅別乗車人員','url':url},{'title':'福岡市 令和6年度高速鉄道事業決算','url':financeUrl},{'title':'福岡県 第5回北部九州圏PT調査','url':'https://www.pref.fukuoka.lg.jp/contents/hokubukyusyuken-persontrip5.html'},{'title':'福岡県 令和2年国勢調査','url':'https://data.bodik.jp/dataset/400009_kokuseichousa_r02'}])
# 2020 Census, Table 6-1: commuters/students by place of residence and workplace/school.
# Municipality totals and ward detail are separate; never sum both levels.
commuteSource='https://www.city.fukuoka.lg.jp/soki/tokeichosa/shisei/toukei/kokusei/R2kokuchou_kouhyou/documents/R2kokuchou_3_6_1.xlsx'
ws3=openpyxl.load_workbook(RAW/'commute-out.xlsx',data_only=True).active
od=[]
for row in range(10,ws3.max_row+1):
 origin=str(ws3.cell(row,1).value or '')
 if ws3.cell(row,2).value!='0_総数' or not origin.startswith('4013'):continue
 oc,on=origin.split('_',1)
 for col in range(3,ws3.max_column+1):
  dest=str(ws3.cell(8,col).value or '')
  if '_' not in dest:continue
  dc,dn=dest.split('_',1)
  level='municipality' if oc=='40130' and dc in names else 'ward' if oc!='40130' and dc.startswith('4013') and dc!='40130' else None
  if not level:continue
  v=ws3.cell(row,col).value
  od.append(dict(origin=on,destination=dn,origin_code=oc,destination_code=dc,people=v if isinstance(v,(int,float)) else None,raw=v,level=level,year=2020,unit='通勤者・通学者／人（移動回数ではない）',source_url=commuteSource,source_cell=ws3.cell(row,col).coordinate))
analysis['commute_od']=od
analysis['sources'].append({'title':'令和2年国勢調査・福岡市 通勤通学OD（第6-1表）','url':commuteSource})
write('data/analysis.json',analysis)
resources=[]
for title,match,url,eligible,note in [('筑後川温泉病院 患者送迎','筑後川温泉病院','https://onsen-hp.or.jp/news/sougeimuryo/','ひとりで乗降できる通院患者','うきは市・朝倉市杷木。前日までに電話予約。'),('福西会南病院 外来送迎','福西会南病院','https://www.fukuseikai-minami.com/contact/','外来再来患者で公共交通利用が困難・不便な方。自立歩行または家族介助。','病院南側を基本に個別調整。受付に事前申込。')]:
 p=next((p for p in places if match in p['name']),None)
 if not p:continue
 resources.append(dict(id='resource-'+p['id'],name=title,provider=match,category='medical',city=p['city'],address=p['address'],lat=p['lat'],lon=p['lon'],locationType='facility',positionNote='厚生労働省公開施設位置。送迎区域・車両位置ではありません。',coordinatePrecision='施設報告座標・入口未確認',positionSourceUrl=p['sourceUrl'],positionDataAsOf=p['dataAsOf'],serviceStatus='published',eligibleUsers=eligible,generalUse='restricted',generalUseNote='対象患者向け。一般住民への開放は未確認。',reservation='required',reservationNote=note,operatingPeriod='公式案内を確認。運行日は個別確認。',vehicleCount=None,capacityNote='未公表',sourceUrl=url,checkedAt=TODAY,dataAsOf=None,sharingAvailability='unknown',coverageNote=note,destinationId=p['id']))
write('gaps/resources.json',resources)
print({'facilities':len(places),'categories':dict(collections.Counter(p['category'] for p in places)),'shopping':len(shopping),'municipalPopulation':len(pop),'railSegments':len(rail['stations']['features']),'staticBusStops':len(bus['features']),'resources':len(resources),'subwayStations':len(stops)})
