#!/usr/bin/env python3
"""Import the supplied research workbook; retain its source cells and definitions.
Usage: python scripts/import-tourism.py /path/to/research.xlsx
The workbook is read only. No network or guessed missing values.
"""
import hashlib
import json
import re
import sys
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
src = Path(sys.argv[1])
w = openpyxl.load_workbook(src, data_only=True)
S = {s.title: [list(r) for r in s.iter_rows(min_row=6, values_only=True) if any(v is not None for v in r)] for s in w}
codes = {'大分県':'44','由布市':'44213','別府市':'44202','日出町':'44341','宇佐市':'44211','大分市':'44201','国東市':'44214','竹田市':'44208','豊後高田市':'44209'}
sources = [dict(zip(['id','publisher','title','period','coverage','locator','note','url'],r)) for r in S['09_出典台帳']]
sources += [dict(id='P05',publisher='大分県',title='おおいた観光データカタログ',period='2025年の掲載指標／2026-09-13確認',coverage='観光庁系列の宿泊・消費、住民受容度など',locator='トップページ・各分析画面',note='県の毎月宿泊調査とは別系列。個別データの注記を確認。',url='https://oita-tourism-data-catalog.com/'),dict(id='T04',publisher='大分県・各交通事業者',title='大分県公開GTFS（既存交通アプリと共通）',period='2026-09-08取得',coverage='16データの運行予定。全交通の網羅ではない',locator='各事業者のデータと出典は交通画面に表示',note='CC BY 4.0。便数は実乗車人数・座席の空きではない。',url='https://odcs.bodik.jp/440001/2023/03/30/gtfs_bus/')]
source_ids = {s['url']:s['id'] for s in sources}
records=[]

def add(area, period, metric, value, unit, series, method, note, url, sheet, row, category='観光客・宿泊', status='公表値', **extra):
    original_unit=unit
    if url and 'toukei_r06toukei.pdf' in url:
        if '消費' in metric or '消費' in series:
            note=(note or '')+' 市推計は大分県2024年の暫定消費単価を使用。県の確報とは区別。'
        if '宿泊' in metric or '宿泊' in series or '宿泊' in (method or ''):
            note=(note or '')+' 市の宿泊客数は入湯税を基に集計。課税対象外施設を含む全宿泊市場の数値ではない。'
    if url and url.endswith('2273969.pdf') and ('宿泊' in metric):
        unit='人泊'
        note=(note or '')+' 原資料の本文は人泊、表は人と表記。延べ宿泊として扱い、実人数とは区別。'
    record=dict(id=f't-{len(records)+1:04}',area=area,areaCode=codes.get(area),period=period,metric=metric,value=value,unit=unit,originalUnit=original_unit,series=series,method=method,note=note or '',sourceId=source_ids.get(url),sourceUrl=url,category=category,status=('未取得' if value is None else status),verification='共有調査ブックに基づく',provenance={'sheet':sheet,'row':row})
    record.update(extra)
    records.append(record)

for i,r in enumerate(S['01_6市町データ'],6):
    area,period,metric,value,unit,method,note,status,url=r
    series='県宿泊統計' if url and url.endswith('2273969.pdf') else f'{area}・{method}'
    category='消費' if '消費' in metric or '販売額' in metric else '観光客・宿泊'
    if 'GW' in period: continue # Same facility-period is imported from the GW sheet below.
    add(area,period,metric,value,unit,series,method,note,url,'01_6市町データ',i,category,status)
for i,r in enumerate(S['02_県年次'],6):
    series,metric,unit,*_=r
    city=next((name for name in codes if name!='大分県' and metric.startswith(name)),None)
    for year,value in zip([2019,2024,2025],r[3:6]):
        if value is None:continue
        if city and year==2025 and any(o['area']==city and o['period']=='2025年' and o['sourceUrl']==r[8] and o['value']==value for o in records):continue
        add(city or '大分県',f'{year}年',metric,value,unit,'共通基準' if series=='共通基準' else '県宿泊統計',r[7],r[7],r[8],'02_県年次',i,'消費' if '消費' in metric else '観光客・宿泊',column=f'{chr(68+[2019,2024,2025].index(year))}{i}')
for i,r in enumerate(S['03_県内交通手段'],6):
    year,pop,mode,numerator,n,rate,note,url=r
    add('大分県',year,mode,rate,'割合',f'観光実態調査・{pop}',f'{pop} n={n:,}・複数回答',note,url,'03_県内交通手段',i,'交通手段','再集計値' if pop=='国内居住者' else '調査値',numerator=numerator,denominator=n,population=pop)
for i,r in enumerate(S['04_滞在と消費'],6):
    group,period,item,value,unit,method,note,url=r
    add('大分県',period,item,value,unit,group,method,note,url,'04_滞在と消費',i,'消費' if '費用' in group or '消費' in group else '滞在・満足度','調査値')
for i,r in enumerate(S['05_由布2025月次'],6):
    if r[0]=='年間計':continue
    for j,label in enumerate(['日帰り','宿泊','外国人日帰り','外国人宿泊'],1):
        add('由布市','2025年'+r[0],label,r[j],'人','由布市観光動態調査・月次','市集計・外国人は内数','月間値。年次値と足さない。',r[8],'05_由布2025月次',i,'月次・地区',column=f'{chr(65+j)}{i}')
for i,r in enumerate(S['06_別府詳細'],6):
    group,period,item,value,unit,method,note,url=r
    if group in ['客数','市推計消費額','県調査']:continue
    add('別府市',period,item,value,unit,'別府市観光動態調査・'+group,method,note,url,'06_別府詳細',i,'消費' if group=='消費単価' else '月次・地区','計算値' if group=='宿泊割合' else '公表値')
for i,r in enumerate(S['07_GW2026'],6):
    area,place,previous,current,rate,period,note,url=r
    for year,value in [(2025,previous),(2026,current)]:
        add(area,f'{year}年4/29〜5/6',place,value,'人','GW施設入場者数','単一施設・8日間',note,url,'07_GW2026',i,'GW・施設',column=f'{"C" if year==2025 else "D"}{i}',place=place)

# Supplemental official figures. Preserve them as a separate series, never sum with P02.
for metric,value,unit,method in [('延べ宿泊者数',8232000,'人泊','観光庁・宿泊旅行統計調査系列。小規模施設を含む'),('外国人延べ宿泊者数',1505000,'人泊','観光庁・宿泊旅行統計調査系列。外国人は内数'),('住民の観光客受容度',.669,'割合','県公式カタログ掲載の住民満足度指標。設問・標本は原典を確認')]:
    add('大分県','2025年',metric,value,unit,'県公式カタログ掲載値',method,'県の従業員10人以上施設調査とは別系列。市町別・地点別の数値には代用しない。','https://oita-tourism-data-catalog.com/','公式カタログ',None,'住民・県公式',verification='公式カタログ表示を照合',status='掲載値')

g=json.loads((ROOT/'oita-mobility/data/map-data.json').read_text())
anchor_specs=[('yufu-station','由布市','由布院駅前','由布院駅前',None),('yufu-takemoto','由布市','岳本','岳本',None),('beppu-station','別府市','別府駅前','別府駅前',None),('beppu-jigoku','別府市','海地獄前','海地獄前','地獄めぐり'),('beppu-kannawa','別府市','鉄輪','鉄輪温泉',None),('hiji-harmony','日出町','ハーモニーランド','ハーモニーランド','ハーモニーランド'),('hiji-yokoku','日出町','暘谷駅前','暘谷駅前',None),('usa-jingu','宇佐市','宇佐八幡','宇佐八幡',None),('usa-safari','宇佐市','サファリ','サファリ','アフリカンサファリ'),('oita-station','大分市','大分駅前','大分駅前',None),('oita-takasaki','大分市','高崎山','高崎山','うみたまご'),('kunisaki-airport','国東市','大分空港','大分空港',None),('kunisaki-center','国東市','国東','国東',None)]
anchors=[]
for aid,area,label,stop_name,gw in anchor_specs:
    matches=[s for s in g['stops'] if s['name']==stop_name]
    assert matches,stop_name
    s=next((s for s in matches if any(n is not None and n>0 for n in s['departures'])),matches[0])
    anchors.append(dict(id=aid,area=area,areaCode=codes[area],name=label,lat=s['lat'],lon=s['lon'],stopId=s['id'],feedId=g['feeds'][s['feedIndex']]['id'],kind='公開GTFSの停留所位置',sourceUrl=g['feeds'][s['feedIndex']]['sourceUrl'],gwFacility=gw,positionNote='交通を調べる基準点。施設の入口・敷地境界を示すものではありません。'))

out={'schemaVersion':1,'researchAsOf':'2026-09-13','editedAt':'2026-09-13','title':'大分観光・周遊データマップ','sourceWorkbook':{'name':src.name,'sha256':hashlib.sha256(src.read_bytes()).hexdigest(),'sheets':w.sheetnames},'records':records,'sources':sources,'tables':S,'anchors':anchors,'municipalities':[{'code':codes[n],'name':n} for n in ['由布市','別府市','日出町','宇佐市','大分市','国東市']],'transport':{'url':'../oita-mobility/data/map-data.json','sha256':hashlib.sha256((ROOT/'oita-mobility/data/map-data.json').read_bytes()).hexdigest(),'dates':g['meta']['dateLabels'],'retrievedAt':g['meta']['retrievedAt']},'revisionNotes':[{'sourceId':'P02','note':'韓国構成比は概要47.8%、詳細表44.1%。詳細の509,951÷1,156,476は約44.1%。概要と詳細の相違を保持。'},{'sourceId':'P02','note':'延べ宿泊者の単位を本文に合わせ人泊として表示。表・調査ブックの原表記「人」も保持。'},{'sourceId':'P05','note':'県公式カタログ掲載値は独立した系列として追加。県宿泊統計と合算しない。'}]}
target=ROOT/'oita-tourism/data/tourism.json'
target.parent.mkdir(parents=True,exist_ok=True)
target.write_text(json.dumps(out,ensure_ascii=False,indent=2)+'\n')
print(json.dumps({'records':len(records),'sources':len(sources),'anchors':len(anchors),'sheets':len(w.sheetnames)},ensure_ascii=False))
